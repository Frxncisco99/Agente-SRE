from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(r"C:\Agente SRE\Documentacion_Tecnica_Centinel.docx")

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def borders(table, color='D9E2F3'):
    tblPr = table._tbl.tblPr
    b = tblPr.first_child_found_in('w:tblBorders')
    if b is None:
        b = OxmlElement('w:tblBorders'); tblPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge
        el = b.find(qn(tag))
        if el is None: el = OxmlElement(tag); b.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),color)

def set_cell(cell, text, bold=False, color='000000', size=9.5):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(str(text)); r.bold = bold; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
    r.font.name = 'Aptos'; r._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); r._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos')
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = True
    borders(t)
    for i,h in enumerate(headers):
        set_cell(t.rows[0].cells[i], h, True, 'FFFFFF', 9); shade(t.rows[0].cells[i], '17365D')
    for ridx,row in enumerate(rows):
        cells = t.add_row().cells
        for i,val in enumerate(row):
            set_cell(cells[i], val, False, '000000', 9)
            shade(cells[i], 'F4F7FB' if ridx % 2 else 'FFFFFF')
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.space_after = Pt(3); p.add_run(text)
    return p

doc = Document()
sec = doc.sections[0]; sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
styles = doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos'); styles['Normal'].font.size=Pt(10.5); styles['Normal'].paragraph_format.space_after=Pt(6)
for name,size in [('Title',25),('Heading 1',17),('Heading 2',12.5),('Heading 3',11)]:
    s=styles[name]; s.font.name='Aptos Display'; s._element.rPr.rFonts.set(qn('w:ascii'),'Aptos Display'); s._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos Display'); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor(0,0,0); s.paragraph_format.space_before=Pt(12); s.paragraph_format.space_after=Pt(6)

title=doc.add_paragraph(style='Title'); title.alignment=WD_ALIGN_PARAGRAPH.LEFT; title.add_run('Documentación técnica del sistema Centinel')
sub=doc.add_paragraph(); sub.paragraph_format.space_after=Pt(14); r=sub.add_run('Detector de anomalías y alerta temprana de TI'); r.italic=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor(80,90,105)
doc.add_paragraph('Este documento recopila las tecnologías, APIs, componentes y flujo operativo utilizados en el proyecto. Centinel procesa telemetría de empresas, identifica tendencias que pueden anticipar incidentes, descarta comportamientos normales del negocio y envía alertas operativas a Telegram mediante n8n.')

doc.add_heading('Resumen del sistema', level=1)
doc.add_paragraph('El proyecto implementa un Early Warning System para el reto Antes de que suene el teléfono. Su objetivo es avisar antes de que una falla de TI afecte a los clientes, manteniendo una tasa baja de falsas alarmas.')
add_table(doc, ['Capa','Responsabilidad','Tecnología principal'], [
    ['Entrada','Leer métricas minuto a minuto y manejar valores faltantes','Python, CSV'],
    ['Detección','Encontrar tendencias de RAM, disco, latencia y errores','NumPy, regresión lineal'],
    ['Contexto','Distinguir incidentes reales de eventos normales','Reglas offline o LLM'],
    ['Operación','Mostrar estado y controlar la simulación','HTTPServer, HTML, CSS, JavaScript'],
    ['Notificación','Enviar alertas al flujo de automatización y a Telegram','HTTP POST, n8n Cloud, Telegram'],
])

doc.add_heading('Arquitectura y flujo de datos', level=1)
doc.add_paragraph('El flujo principal es causal: cada fila se procesa en orden cronológico y una alerta en el minuto t solo utiliza información disponible hasta ese minuto.')
for x in ['metricas.csv → streaming_simulator.py','streaming_simulator.py → trend_analyzer.py','trend_analyzer.py → llm_context_filter.py','llm_context_filter.py → n8n_webhook_trigger.py','n8n Webhook → nodo Telegram → mensaje al destinatario']:
    add_bullet(doc, x)
doc.add_paragraph('El sistema puede operar por tenant individual, por todos los tenants o por una corrida específica del dataset. El modo offline es reproducible y no necesita red para analizar las anomalías.')

doc.add_heading('Tecnologías utilizadas', level=1)
add_table(doc, ['Tecnología','Uso en el proyecto','Estado'], [
    ['Python 3.10+','Backend, procesamiento, simulación y servidor local','Principal'],
    ['NumPy','Pendiente, regresión OLS y cálculo de R²','Instalada'],
    ['Requests','Peticiones HTTP al webhook de n8n, timeouts y reintentos','Instalada'],
    ['CSV, JSON y Markdown','Métricas, resultados, contexto y logs','Nativo / archivos'],
    ['HTML5, CSS3 y JavaScript vanilla','Interfaz del dashboard','Activo'],
    ['SVG','Gráfica de telemetría en el dashboard','Activo'],
    ['Python HTTPServer','Servidor de localhost:8080 y API local','Activo'],
    ['Google Fonts','IBM Plex Sans, IBM Plex Mono y Outfit','CDN opcional'],
    ['Git y GitHub','Control de versiones y dataset Zikit','Usado'],
    ['PowerShell','Ejecución local en Windows','Usado'],
])

doc.add_heading('APIs y servicios externos', level=1)
add_table(doc, ['API o servicio','Cómo se utiliza','Configuración'], [
    ['Webhook de n8n','Recibe las alertas mediante HTTP POST','N8N_WEBHOOK_URL'],
    ['Telegram Bot API','Entrega el mensaje al usuario final a través del nodo Telegram de n8n','Token del bot y Chat ID en n8n'],
    ['Google Gemini API','Filtro inteligente opcional basado en contexto, métricas y logs','GOOGLE_API_KEY'],
    ['OpenAI API','Alternativa opcional para el filtro inteligente','OPENAI_API_KEY'],
    ['API local del dashboard','Consulta el estado y controla la simulación','GET /api/state y POST /api/control'],
])
doc.add_paragraph('La URL de producción configurada actualmente es:')
p=doc.add_paragraph(); r=p.add_run('https://frxncisco99.app.n8n.cloud/webhook/ews-alert'); r.font.name='Consolas'; r.font.size=Pt(9); r.font.color.rgb=RGBColor(31,78,121)
doc.add_paragraph('El proyecto no llama directamente a Telegram. El flujo es Python → n8n → Telegram Bot API.')

doc.add_heading('Detección de anomalías', level=1)
add_table(doc, ['Patrón','Señal analizada','Interpretación'], [
    ['memory_leak','RAM sube de forma sostenida y lineal','Posible agotamiento de memoria'],
    ['disk_leak','Disco aumenta de manera sostenida','Riesgo para respaldos y operación'],
    ['resource_exhaustion','Latencia sube mientras CPU permanece estable','Pool agotado, bloqueo o deadlock'],
    ['error_burst','Errores aumentan rápidamente','Posible degradación de aplicación o red'],
])
doc.add_paragraph('Las ventanas configuradas son de 15, 30 y 60 minutos. Se utilizan pendiente, consistencia de la tendencia mediante R², score compuesto y un cooldown para evitar alertas repetidas.')

doc.add_heading('Filtro de contexto y falsas alarmas', level=1)
doc.add_paragraph('El sistema consulta contexto.md para conocer horarios, respaldos, promociones, cierres de nómina y otros eventos normales. También puede extraer logs anteriores al timestamp de la anomalía. Una alerta solo se envía cuando el filtro la clasifica como incidente real y alcanza la confianza mínima configurada, actualmente 0.75.')
add_bullet(doc, 'Modo offline: reglas deterministas, sin API externa.')
add_bullet(doc, 'Modo Gemini: análisis contextual con Google Gemini.')
add_bullet(doc, 'Modo OpenAI: alternativa mediante la API de OpenAI.')

doc.add_heading('Dashboard', level=1)
doc.add_paragraph('El dashboard se sirve con dashboard_server.py y se abre en http://localhost:8080. La interfaz muestra el tenant activo, estado general, puntos procesados, anomalías, incidentes reales, señuelos filtrados, métricas actuales, gráfica de latencia y feed de alertas.')
add_table(doc, ['Endpoint','Método','Función'], [
    ['/api/state','GET','Devuelve métricas, historial, estadísticas y anomalías recientes.'],
    ['/api/control','POST','Pausa, reanuda, reinicia, cambia tenant o ajusta la velocidad.'],
])
doc.add_paragraph('El envío de alertas del dashboard se ejecuta en segundo plano para que un timeout o una falla temporal de n8n no congele la visualización.')

doc.add_heading('Archivos principales', level=1)
add_table(doc, ['Archivo','Responsabilidad'], [
    ['config.py','Rutas, tenants, umbrales, proveedor LLM y webhook de n8n.'],
    ['models.py','Modelos MetricPoint, DetectedAnomaly, LLMVerdict y VerifiedAlert.'],
    ['streaming_simulator.py','Lectura causal de métricas e imputación limitada.'],
    ['trend_analyzer.py','Detección matemática de tendencias e incidentes.'],
    ['llm_context_filter.py','Filtro por contexto, logs y proveedor LLM.'],
    ['n8n_webhook_trigger.py','Construcción del payload, POST, timeout y reintentos.'],
    ['main.py','Orquestación completa por tenant o corrida.'],
    ['dashboard_server.py','Servidor local y simulación visual.'],
    ['dashboard.html','Interfaz visual Centinel del dashboard.'],
    ['evaluate_competition.py','Evaluación inicial contra etiquetas de referencia.'],
])

doc.add_heading('Comandos principales', level=1)
for cmd in ['python main.py --tenant gracko --offline','python main.py --all --offline','python main.py --corrida 48 --tenant gracko --offline','python dashboard_server.py 8080']:
    p=doc.add_paragraph(style='List Bullet'); r=p.add_run(cmd); r.font.name='Consolas'; r.font.size=Pt(9)

doc.add_heading('Siguiente integración prevista', level=1)
doc.add_paragraph('La siguiente fase consiste en agregar una nueva empresa alimentada por dos computadoras reales mediante dispositivos LoRa. Para ello se incorporará un adaptador que transforme los datos recibidos en el formato MetricPoint, manteniendo el mismo detector, filtro de contexto, dashboard y flujo n8n → Telegram.')
add_bullet(doc, 'Identificar el transporte usado por LoRa: serial, USB, MQTT, HTTP u otro.')
add_bullet(doc, 'Asignar un identificador a cada computadora y a cada métrica.')
add_bullet(doc, 'Agregar nombre, giro, horarios y eventos normales de la nueva empresa.')
add_bullet(doc, 'Registrar el tenant en el dashboard y conservar la trazabilidad de cada dispositivo.')

doc.add_heading('Riesgos y límites actuales', level=1)
add_bullet(doc, 'La evaluación automática de las 54 corridas todavía requiere mayor consolidación.')
add_bullet(doc, 'El rendimiento depende de que n8n esté activo y accesible por Internet.')
add_bullet(doc, 'Gemini y OpenAI son opcionales; el modo offline es el más reproducible.')
add_bullet(doc, 'La integración LoRa aún depende de conocer el formato y transporte de los scripts reales.')
add_bullet(doc, 'Los resultados de competencia deben validarse por corrida, tenant, precisión, recall y anticipación.')

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; footer.add_run('Centinel · Documentación técnica').font.size=Pt(8)
doc.core_properties.title='Documentación técnica del sistema Centinel'; doc.core_properties.subject='Tecnologías, APIs y arquitectura'; doc.core_properties.author='Equipo Centinel'
doc.save(OUT)
print(OUT)
